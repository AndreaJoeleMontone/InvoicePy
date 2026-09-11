from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


HEADERS = [
    "File",
    "Fornitore",
    "Numero fattura",
    "Data emissione",
    "Scadenza",
    "Importo",
    "Riconoscimento %",
    "Stato",
    "Campi mancanti",
]


def export_to_excel(invoices, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Fatture"

    worksheet.append(HEADERS)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for invoice in invoices:
        missing_fields = ", ".join(invoice.get("campi_mancanti", []))

        worksheet.append([
            invoice.get("file"),
            invoice.get("fornitore"),
            invoice.get("numero_fattura"),
            invoice.get("data_emissione"),
            invoice.get("scadenza"),
            invoice.get("importo"),
            invoice.get("confidence"),
            invoice.get("status"),
            missing_fields,
        ])

    for row in worksheet.iter_rows(min_row=2):
        amount_cell = row[5]
        if isinstance(amount_cell.value, (int, float)):
            amount_cell.number_format = '#,##0.00 [$€-it-IT]'

        confidence_cell = row[6]
        if isinstance(confidence_cell.value, (int, float)):
            confidence_cell.number_format = '0"%"'

    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )

        column_letter = get_column_letter(column_cells[0].column)
        worksheet.column_dimensions[column_letter].width = min(max_length + 3, 40)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    workbook.save(output_path)
